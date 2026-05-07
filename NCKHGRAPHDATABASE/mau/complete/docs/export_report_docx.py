from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
NUMBERED_RE = re.compile(r"^(\d+)\.\s+(.*)$")


def strip_inline_markdown(text: str) -> str:
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^\)]+\)", r"\1", text)
    return text.strip()


def set_cell_text(cell, text: str) -> None:
    cell.text = strip_inline_markdown(text)
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(13)


def apply_base_style(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(2)

    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(13)


def add_page_number(section) -> None:
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = 1

    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def add_paragraph(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(strip_inline_markdown(text))
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(13)


def parse_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    table_lines: list[str] = []
    while index < len(lines) and lines[index].strip().startswith("|"):
        table_lines.append(lines[index].strip())
        index += 1

    rows: list[list[str]] = []
    for line_number, line in enumerate(table_lines):
        cells = [cell.strip() for cell in line.strip("|").split("|")]
        if line_number == 1 and all(set(cell) <= {"-", ":"} for cell in cells):
            continue
        rows.append(cells)
    return rows, index


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for idx, cell_text in enumerate(rows[0]):
        set_cell_text(header_cells[idx], cell_text)

    for row_values in rows[1:]:
        row_cells = table.add_row().cells
        for idx, cell_text in enumerate(row_values):
            set_cell_text(row_cells[idx], cell_text)

    document.add_paragraph()


def add_code_block(document: Document, block_lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(1)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run("\n".join(block_lines))
    run.font.name = "Courier New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    run.font.size = Pt(10.5)


def convert_markdown_to_docx(input_path: Path, output_path: Path) -> None:
    document = Document()
    apply_base_style(document)
    add_page_number(document.sections[0])

    lines = input_path.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code_block = False
    code_lines: list[str] = []

    while index < len(lines):
        raw_line = lines[index]
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code_block:
                add_code_block(document, code_lines)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            index += 1
            continue

        if in_code_block:
            code_lines.append(raw_line)
            index += 1
            continue

        if not stripped:
            index += 1
            continue

        if stripped == "---":
            document.add_page_break()
            index += 1
            continue

        heading_match = HEADING_RE.match(stripped)
        if heading_match:
            level = min(len(heading_match.group(1)), 4)
            text = strip_inline_markdown(heading_match.group(2))
            paragraph = document.add_heading(text, level=level)
            paragraph.paragraph_format.space_after = Pt(6)
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            index += 1
            continue

        if stripped.startswith("|"):
            rows, next_index = parse_table(lines, index)
            add_table(document, rows)
            index = next_index
            continue

        if stripped.startswith(">"):
            add_paragraph(document, stripped.lstrip("> "), style="Intense Quote")
            index += 1
            continue

        if stripped.startswith("- "):
            add_paragraph(document, stripped[2:], style="List Bullet")
            index += 1
            continue

        numbered_match = NUMBERED_RE.match(stripped)
        if numbered_match:
            add_paragraph(document, numbered_match.group(2), style="List Number")
            index += 1
            continue

        add_paragraph(document, stripped)
        index += 1

    document.save(output_path)


def main() -> int:
    if len(sys.argv) != 3:
        print("Usage: python export_report_docx.py <input.md> <output.docx>")
        return 1

    input_path = Path(sys.argv[1]).resolve()
    output_path = Path(sys.argv[2]).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    convert_markdown_to_docx(input_path, output_path)
    print(f"Created DOCX: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())